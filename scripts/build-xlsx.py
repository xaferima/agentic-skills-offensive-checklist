#!/usr/bin/env python3
# build-xlsx.py — build a tracking spreadsheet from checklist/items.json.
# Sheet "Checklist": one row per item with a Status dropdown.
# Sheet "Dashboard": per-section and per-status live counts.
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CHK  = os.path.join(BASE, "checklist")
items = json.load(open(os.path.join(CHK, "items.json"), encoding="utf-8"))

STATUSES = ["Pendiente", "En progreso", "Confirmado", "No aplica", "Falso positivo"]

wb = Workbook()
ws = wb.active
ws.title = "Checklist"

HDR = ["Nº", "Sección", "Source ID", "Título", "Impacto", "Severidad", "Target", "Estado", "Notas / Evidencia"]
hdr_fill = PatternFill("solid", fgColor="1F3864")
hdr_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.append(HDR)
for c in range(1, len(HDR) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hdr_fill
    cell.font = hdr_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border

for it in items:
    ws.append([it["num"], it["sec_title"], it["id"], it["title"],
               it["impact"], it["severity"], it["target"], "Pendiente", ""])

last = ws.max_row
for r in range(2, last + 1):
    for c in range(1, len(HDR) + 1):
        ws.cell(row=r, column=c).border = border
        ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=(c in (4, 9)))

# Status dropdown
dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
dv.error = "Elige un estado de la lista"
dv.errorTitle = "Estado inválido"
ws.add_data_validation(dv)
dv.add(f"H2:H{last}")

# Severity color coding (visual only)
sev_fill = {
    "critical": PatternFill("solid", fgColor="F8CBAD"),
    "high":     PatternFill("solid", fgColor="FFD966"),
    "medium":   PatternFill("solid", fgColor="FFF2CC"),
    "low":      PatternFill("solid", fgColor="E2EFDA"),
}
for r in range(2, last + 1):
    sev = str(ws.cell(row=r, column=6).value or "").strip().lower()
    if sev in sev_fill:
        ws.cell(row=r, column=6).fill = sev_fill[sev]

widths = [7, 34, 11, 46, 13, 10, 11, 16, 26]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(HDR))}{last}"

# ---------- Dashboard ----------
db = wb.create_sheet("Dashboard")
db["A1"] = "Dashboard — Agentic Skills Offensive Checklist"
db["A1"].font = Font(bold=True, size=14, color="1F3864")
db["A2"] = f"Total de ítems: {len(items)}"
db["A2"].font = Font(bold=True)

# per-status
db["A4"] = "Por estado"; db["A4"].font = Font(bold=True)
row = 5
for i, st in enumerate(STATUSES):
    db.cell(row=row, column=1, value=st)
    db.cell(row=row, column=2, value=f'=COUNTIF(Checklist!$H$2:$H${last},"{st}")')
    row += 1

# per-section
db["D4"] = "Por sección"; db["D4"].font = Font(bold=True)
db.cell(row=5, column=4, value="Sección")
db.cell(row=5, column=5, value="Ítems")
db.cell(row=5, column=6, value="Confirmados")
secs = []
for it in items:
    if it["sec_title"] not in secs:
        secs.append(it["sec_title"])
r = 6
for s in secs:
    safe = s.replace('"', '""')
    db.cell(row=r, column=4, value=s)
    db.cell(row=r, column=5, value=f'=COUNTIF(Checklist!$B$2:$B${last},"{safe}")')
    db.cell(row=r, column=6, value=f'=COUNTIFS(Checklist!$B$2:$B${last},"{safe}",Checklist!$H$2:$H${last},"Confirmado")')
    r += 1
db.column_dimensions["A"].width = 20
db.column_dimensions["B"].width = 10
db.column_dimensions["D"].width = 46
db.column_dimensions["E"].width = 8
db.column_dimensions["F"].width = 12

out = os.path.join(CHK, "agentic-skills-offensive-checklist.xlsx")
wb.save(out)
print(f"wrote {out} rows={len(items)} statuses={len(STATUSES)} sections={len(secs)}")
