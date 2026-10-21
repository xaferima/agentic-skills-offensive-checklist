#!/usr/bin/env python3
"""check-xlsx.py — integrity check for the committed tracking workbook.

Verifies the workbook opens, has the expected sheets, 65 data rows and the
status dropdown on H2:H66. Exit 1 on any problem (used by scripts/ci.sh).
"""
import os
import sys

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH = os.path.join(BASE, "checklist", "agentic-skills-offensive-checklist.xlsx")
EXPECTED_ROWS = 65

errs = []
try:
    wb = load_workbook(PATH)
except Exception as exc:  # noqa: BLE001
    print(f"XLSX FAIL: cannot open {PATH}: {exc}")
    sys.exit(1)

if set(wb.sheetnames) != {"Checklist", "Dashboard"}:
    errs.append(f"sheets={wb.sheetnames}")

ws = wb["Checklist"] if "Checklist" in wb.sheetnames else None
rows = 0
if ws is not None:
    rows = ws.max_row - 1
    if rows != EXPECTED_ROWS:
        errs.append(f"data rows={rows} (expected {EXPECTED_ROWS})")
    dvs = [(dv.type, str(dv.sqref)) for dv in ws.data_validations.dataValidation]
    if not any(kind == "list" and "H2:H66" in rng for kind, rng in dvs):
        errs.append(f"status dropdown missing on H2:H66 (found {dvs})")

if errs:
    print("XLSX FAIL: " + "; ".join(errs))
    sys.exit(1)

print(f"xlsx ok: sheets={wb.sheetnames} rows={rows} dropdown=H2:H66")
